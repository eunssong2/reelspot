"""브랜치 실행 캡처를 우선 사용해 UI 설계서와 추적 자료를 함께 생성한다."""
import base64, hashlib, html, json
from pathlib import Path
from xml.sax.saxutils import escape
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from design_model import ROOT, SCREENS

E = html.escape
FONT = 'Noto Sans CJK KR'
RELATED = {'S0':None,'S02':'UI08','S02-empty':'UI08','S02-oneday':'UI08','S02-tabs':'UI08',
           'M4':'UI08','M1':'UI12','S05-empty':'UI12','S03-detail':'UI10',
           'V1-video':'UI18','M3':'UI14','S04-empty':'UI13'}


def load_source():
    manifest=json.loads((ROOT/'UI/프로토타입/출처.json').read_text())
    for shot in manifest['screens']:
        p=ROOT/'UI/프로토타입'/shot['image']
        if hashlib.sha256(p.read_bytes()).hexdigest()!=shot['sha256']:
            raise ValueError('프로토타입 캡처 SHA256 불일치: '+p.name)
    return manifest


def build(screen_elements, svg_for, ppt_shape, sheet):
    manifest=load_source()
    shots=manifest['screens']; primary={s['ui']:s for s in shots if s['ui']}
    commit=manifest['commit']; short=commit[:7]
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    def text(slide,x,y,w,h,value,size=18,color='191F28',bold=False):
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        box.text_frame.word_wrap=True
        box.text_frame.margin_left=box.text_frame.margin_right=0
        box.text_frame.margin_top=box.text_frame.margin_bottom=0
        for i,line in enumerate(value.split('\n')):
            p=box.text_frame.paragraphs[0] if i==0 else box.text_frame.add_paragraph()
            p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold
            p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(5)
        return box
    def base(title,sub):
        s=prs.slides.add_slide(prs.slide_layouts[6])
        sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,Inches(.85))
        sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string('202A35');sh.line.fill.background()
        text(s,.5,.16,15,.5,title,25,'FFFFFF',True)
        text(s,.5,8.57,15,.3,sub+' · 2026-10-26',10,'5F6B78')
        return s
    cover=base('At that place · UI 설계 · 프로토타입 화면 반영',f'feature/ui-prototype @ {short}')
    text(cover,.7,1.25,10,1.2,'일정 요약에서 장소별 기록까지\n브랜치 화면을 설계 기준으로 반영합니다.',32,bold=True)
    text(cover,.7,3.15,9.3,3.4,
         f'원격 브랜치: feature/ui-prototype\n기준 커밋: {commit}\n실행 캡처 {len(shots)}개 · 기존 UI ID {len(SCREENS)}개 유지\n7개 대응 화면은 캡처로 교체하고 관련 상태 12개를 추가했습니다.\n그 외 26개 화면은 기존 명세 기반 설계안으로 유지합니다.\n캡처 이미지는 원본 그대로이며 설명·계약 텍스트는 편집 가능합니다.',19)
    text(cover,.7,7,9.4,.9,'제품 정책: 기능명세서 v2.1 · 시스템/API/DB 설계 v0.3.1\n화면 변경과 서비스 구현·검수 완료를 구분합니다.',17,'B8433F',True)
    cover.shapes.add_picture(str(ROOT/'UI/프로토타입/S01.png'),Inches(11.6),Inches(1.1),height=Inches(7.2))
    flow=base('화면 흐름과 원본 위치','브랜치 S 번호와 설계 UI 번호를 함께 표기')
    text(flow,.8,1.35,14.5,2.5,
         'S0 데모 선택 → S01 일정 요약(UI08) → S02 날짜별 장소(UI08 상세)\nS02 → S03 장소 추가(UI10) / S05 순서 편집(UI12) / M1 방문 시간\nS02 → S04 장소 기록(UI13) → 사진별·사람별(UI09) → V1 감상(UI18)\nS04 → M2 기록 추가(UI14) → M3 촬영·앨범 선택',23)
    text(flow,.8,4.25,14.5,2.2,
         '화면 원본: src/prototype/screens/*.tsx · 관련 시트/팝업: components/ 및 src/demo/\nUI/프로토타입/출처.json: 커밋·소스 해시·화면별 실행 경로·캡처 해시\nUI/프로토타입/demo-app-spec.txt · feature-spec.txt: 해당 커밋의 화면 명세 사본\n재생성은 저장된 캡처를 사용하며 원격 브랜치·Expo 실행이 필요하지 않습니다.',20)
    text(flow,.8,7.05,14.3,.9,'UI09 캡처는 한 방문의 사람별 보기입니다. 일정 전체·미지정·기간 밖 조회는 제품 계약으로 남습니다.',17,'B8433F',True)
    differences=[
        ['날짜 추가','S01 마지막 날 다음 날짜 추가·최대30일','UI06 새 일정 생성과 구분·제품 기간 최대366일'],
        ['메모','S05 장소 메모 입력100자 남음','날짜 공통 메모 제외'],
        ['방문 시간','M1 휠·1분 단위','기존 제품 설계의5분 단위와 차이'],
        ['기록 정렬','게시물 시각 최신순·보기 전환은 진입마다 초기화','촬영 시각 오름차순·보기 방식 기기 기억'],
        ['파일 선택','M2 최대10개·실제 업로드 연결 지점','방문당 작성자 사진5/영상2·다중선택 최대120'],
        ['서버·권한·기한','예시 데이터·로컬 저장·예시 지도','권한·한도·검사·보관·버전·실패 복구 별도 구현'],
    ]
    diff=base('화면과 제품 계약의 차이','화면을 반영하면서 기존 정책을 임의로 바꾸지 않습니다.')
    y=1.25
    for name,actual,contract in differences:
        text(diff,.75,y,2,.7,name,18,'B8433F',True)
        text(diff,2.9,y,5.8,.8,actual,17)
        text(diff,9,y,6,.8,contract,17)
        y+=1.1
    text(diff,.8,7.95,14.5,.4,'이 자료의 캡처는 웹 실행 결과입니다. Android APK·실서비스·외부 연동 검수 결과로 해석하지 않습니다.',13,'5F6B78')

    articles=[];screen_rows=[];components=[];states=[];apis=[];data=[];page_ids=[]
    nav=[]
    def article(id,title,kind,img,notes,related='',extra=''):
        nav.append(f'<a href="#{E(id)}">{E(id)} {E(title)}</a>');page_ids.append(id)
        articles.append(f'<article class="screen" id="{E(id)}"><h2>{E(id)} · {E(title)}</h2>'
                        f'<p class="meta">{E(kind)} · {E(related)}</p><div class="pair"><figure>'
                        f'<img src="{E(img)}" alt="{E(title)}"><figcaption>{E(kind)}</figcaption></figure>'
                        f'<div class="notes">'+''.join(f'<p>{E(n)}</p>' for n in notes)+extra+'</div></div></article>')
    def capture_page(shot,ui_spec=None):
        id=ui_spec['id'] if ui_spec else shot['id']
        title=shot['title'];related=shot['ui'] or RELATED.get(shot['id']) or '데모 전용'
        s=base(id+'  '+title,f'브랜치 실행 캡처 · {shot["id"]} · {short}')
        s.shapes.add_picture(str(ROOT/'UI/프로토타입'/shot['image']),Inches(.65),Inches(1.05),height=Inches(7.2))
        text(s,4.3,1.15,10.9,.5,shot['id']+' · '+shot['actions'],22,bold=True)
        y=2
        for n in shot['notes']:
            text(s,4.3,y,10.8,.9,n,18);y+=1.05
        text(s,4.3,5.5,10.8,.35,'출처와 적용 범위',16,'B8433F',True)
        text(s,4.3,5.95,10.8,1.5,shot['source']+'\n'+manifest['environment']+'\n'+manifest['scope'],14,'5F6B78')
        if ui_spec:
            text(s,4.3,7.52,10.8,.8,'제품 계약: '+' / '.join(ui_spec['features'])+' · '+ui_spec['role']+'\nAPI: '+', '.join(ui_spec['api']),12,'5F6B78')
        s.notes_slide.notes_text_frame.text='실행경로: '+shot['actions']+'\n브랜치: '+manifest['branch']+'\n커밋: '+commit+'\n'+ '\n'.join(shot['notes'])
        if ui_spec:
            s.notes_slide.notes_text_frame.text+='\n제품 구현 계약(데모에서 미구현 가능):\n'+'\n'.join(ui_spec['notes'])+'\n상태: '+ui_spec['states']
        extra=f'<h3>원본과 실행 경로</h3><p>{E(shot["source"])}</p><p>{E(shot["actions"])}</p><p class="small">{E(shot["route"])}</p>'
        if ui_spec:
            # 제품 계약은 접어서 제공한다. 인쇄 시에는 핵심 상태와 API를 함께 표시한다.
            extra+=f'<details><summary>제품 구현 계약 · {E(ui_spec["title"])}</summary>'+''.join(f'<p>{E(n)}</p>' for n in ui_spec['notes'])+'</details>'
            extra+=f'<h3>제품 상태·API</h3><p class="small">{E(ui_spec["states"])}</p><p class="small">{E(", ".join(ui_spec["api"]))}</p>'
        article(id,title,'브랜치 실행 캡처 · '+shot['id']+' @ '+short+' · 웹','프로토타입/'+shot['image'],shot['notes'],related,extra)
    for spec in SCREENS:
        id=spec['id'];filename=id+'_'+spec['title'].replace(' ','_')+'.svg'
        path=ROOT/'UI/SVG'/filename
        shot=primary.get(id)
        if shot:
            # 기존 SVG 경로도 최신 캡처를 담아, 외부의 기존 링크가 초안으로 돌아가지 않게 한다.
            b64=base64.b64encode((ROOT/'UI/프로토타입'/shot['image']).read_bytes()).decode()
            path.write_text(f'<svg xmlns="http://www.w3.org/2000/svg" width="390" height="844" viewBox="0 0 390 844" role="img"><title>{escape(shot["title"])}</title><desc>feature/ui-prototype @ {commit} · 웹 실행 캡처 · 이미지 편집은 앱 소스에서 수행</desc><image width="390" height="844" href="data:image/png;base64,{b64}"/></svg>')
            capture_page(shot,spec)
            title=shot['title'];kind='브랜치 실행 캡처';image='프로토타입/'+shot['image']
            display=dict(spec, title=title, rows=[['prototype',t] for t in shot.get('text','').splitlines()],
                         contract_title=spec['title'],contract_rows=spec['rows'],
                         visual_source=kind,prototype=shot,source_commit=commit)
            for i,t in enumerate(shot.get('text','').splitlines(),1):components.append([id,f'{id}-P{i:02}','실행 화면 문구',t,shot['actions'],'데모 표시·서비스 권한 별도',shot['id']])
        else:
            path.write_text(svg_for(screen_elements(spec)))
            title=spec['title'];kind='명세 기반 설계안 · 기존 화면 유지';image='미리보기/'+path.stem+'.png'
            s=base(id+'  '+title,kind)
            for e in screen_elements(spec):ppt_shape(s,e,.65,1.05,3.33/390)
            text(s,4.4,1.2,10.7,.5,'대상: '+spec['role'],21,bold=True)
            y=1.95
            step=4.65/len(spec['notes'])
            for n in spec['notes']:
                text(s,4.4,y,10.6,step-.1,n,16);y+=step
            text(s,4.4,6.75,10.6,.9,spec['states'],14,'5F6B78')
            text(s,4.4,7.9,10.6,.4,'API: '+', '.join(spec['api']),11,'5F6B78')
            s.notes_slide.notes_text_frame.text='\n'.join(spec['notes'])+'\n'+spec['states']
            article(id,title,kind,image,spec['notes'],' / '.join(spec['features']),f'<h3>상태·API</h3><p>{E(spec["states"])}</p><p class="small">{E(", ".join(spec["api"]))}</p><p class="small"><a href="SVG/{E(filename)}">편집 가능한 SVG 원본</a></p>')
            display=dict(spec,visual_source=kind)
            for i,(k,t) in enumerate(spec['rows'],1):components.append([id,f'{id}-C{i:02}',k,t,'\n'.join(spec['notes']),spec['role'],spec['states']])
        data.append(display)
        screen_rows.append([id,title,' '.join(spec['features']),spec['role'],kind,image,shot['id'] if shot else '',shot['source'] if shot else '도구/design_model.py',spec['states'],', '.join(spec['api'])])
        states.append([id,kind,spec['states'],'제품 요구: 입력 보존·오류/빈 결과 구분·권한/버전 재검증 · 데모 구현과 구분'])
        apis.extend([[id,a,' '.join(spec['features']),'제품 구현 계약 · 데모 연동 검수 아님'] for a in spec['api']])
        for extra in shots:
            if not extra['ui'] and RELATED.get(extra['id'])==id:capture_page(extra)
    capture_page(next(s for s in shots if s['id']=='S0'))
    prs.save(ROOT/'UI/05 UI 설계.pptx')
    wb=Workbook();wb.remove(wb.active)
    sheet(wb,'화면목록',['화면ID','화면명','기능','권한','화면출처','대표이미지','브랜치화면','원본소스','제품상태','API'],screen_rows)
    sheet(wb,'구성요소명세',['화면ID','요소ID','유형','표시문구','동작·검증·전환','권한·범위','상태·원본'],components)
    sheet(wb,'화면API연결',['화면ID','operationId','기능','범위'],apis)
    sheet(wb,'화면상태',['화면ID','화면출처','제품별상태','복구규칙'],states)
    source_rows=[[s['id'],s['title'],s['ui'] or RELATED.get(s['id']) or '데모 전용',s['source'],s['route'],s['actions'],s['captured_at'],s['sha256']] for s in shots]
    sheet(wb,'프로토타입대응',['브랜치화면','화면명','설계UI','소스','라우트','실행경로','캡처시각','캡처SHA256'],source_rows)
    sheet(wb,'정책차이',['항목','브랜치데모','기존제품계약'],differences,[25,75,75])
    flows=[['UI05','UI08/S01','일정 선택','제품 진입'],['S0','S01','예시 일정 선택','데모 전용'],['S01','S02','날짜 카드 선택','tripId/itineraryId/dayId'],['S01','M4','배경 사진 변경','앨범 선택'],['S02','S03','장소 추가','선택 날짜'],['S03','S02','추가 성공','마지막 순서·시간 미정·추가됨'],['S02','S05','편집','빈 날짜에도 제공'],['S02/S05','M1','시간 수정','S05에서는 변경 저장 때 반영'],['S05','S02','변경 저장','순서·시각·제목'],['S02','S04','장소 카드','방문별 기록'],['S04','V1','파일 탭','사진 확대·영상 재생'],['S04','M2','올리기','업로드 연결 지점'],['M2','M3','추가 타일','촬영·앨범 선택']]
    # 브랜치 전환과 기존 제품 계약을 구분해 모두 보존한다.
    flows+=json.loads((ROOT/'UI/제품 화면전환.json').read_text())
    sheet(wb,'화면전환',['출발','도착','이벤트','범위·전달'],flows,[25,25,35,80])
    sheet(wb,'디자인토큰',['토큰','브랜치값','원본'],[
        ['primary','#F1706B','src/theme/theme.ts'],['primaryPressed','#DC5F5A','src/theme/theme.ts'],['primarySoft','#FDEDEB','src/theme/theme.ts'],['text','#191F28','src/theme/theme.ts'],['surface','#F2F4F6','src/theme/theme.ts'],['pageBg','#F9FAFB','src/prototype/styles/tokens.ts'],['간격','4px 단위·좌우20px','src/theme/theme.ts'],['큰 버튼','최소56px·radius14','src/prototype/styles/tokens.ts'],['카드/행','radius20/16','src/prototype/styles/tokens.ts'],['캡처','390×844 CSS px·DPR2','Playwright Chromium 웹'],['브랜치',manifest['branch'],commit]], [30,48,85])
    wb.save(ROOT/'UI/05 화면 명세.xlsx')
    (ROOT/'UI/화면데이터.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
    css="""*{box-sizing:border-box}body{font:15px/1.8 'Noto Sans CJK KR',sans-serif;color:#191F28;background:#f4f5f7;margin:0}a{color:#B8433F}aside{position:fixed;inset:0 auto 0 0;width:260px;background:white;overflow:auto;padding:24px}aside a{display:block;font-size:12px;padding:6px;text-decoration:none}main{margin-left:260px}header{background:#202A35;color:white;padding:40px}header a{color:white}.screen,.intro{background:white;margin:28px;padding:28px;border-radius:18px;scroll-margin:20px}h2{font-size:25px;margin:0 0 8px}.meta{color:#B8433F;font-size:13px}.pair{display:flex;gap:32px;align-items:flex-start}figure{width:260px;flex-shrink:0;margin:0}figure img{width:100%;border:1px solid #dce1e6}figcaption,.small{font-size:11px;overflow-wrap:anywhere}.notes{max-width:760px}.notes p{margin:0 0 16px}.notes h3{font-size:15px;margin:15px 0 8px}details{background:#f2f4f6;padding:12px;margin-bottom:16px}details p{font-size:13px}table{border-collapse:collapse;width:100%;font-size:12px}th,td{border:1px solid #dce1e6;padding:8px;text-align:left}th{background:#f2f4f6}@media(max-width:1100px){aside{display:none}main{margin:0}.pair{flex-wrap:wrap}}@media print{@page{size:A4 landscape;margin:10mm}aside,header,.intro{display:none}main{margin:0}.screen{height:187mm;margin:0;padding:0;border-radius:0;break-after:page;break-inside:avoid}.screen:last-child{break-after:auto}.pair{flex-wrap:nowrap;gap:24px}figure{width:232px}figure img{height:502px;object-fit:contain}.notes{font-size:12px;line-height:1.7}.notes p{margin-bottom:11px}.notes h3{font-size:12px}.small{font-size:9px}details{display:none}body{background:white}h2{font-size:21px}.meta{margin:4px 0 12px}}"""
    table='<table><tr><th>항목</th><th>브랜치 데모</th><th>제품 구현 계약</th></tr>'+''.join('<tr>'+''.join('<td>'+E(c)+'</td>' for c in row)+'</tr>' for row in differences)+'</table>'
    css+='@media print{.intro{display:block;height:187mm;margin:0;padding:0;break-after:page;font-size:12px}.intro table{font-size:10px}.intro td,.intro th{padding:6px}}'
    intro=f'<section class="intro"><h2>화면 적용 기준</h2><p>feature/ui-prototype @ <b>{short}</b> · 2026-10-26 · 브랜치 실행 캡처 19개. 기존 33개 UI 중 7개를 교체하고 관련 상태 12개를 추가했습니다. 나머지 26개는 명세 기반 설계안입니다.</p><p>{E(manifest["environment"])}<br>{E(manifest["scope"])}</p><p><a href="프로토타입/출처.json">커밋·소스·캡처 해시와 경로</a> · <a href="프로토타입/demo-app-spec.txt">브랜치 화면 명세 원문</a> · <a href="프로토타입/feature-spec.txt">브랜치 기능 명세 원문</a></p><h3>제품 정책과 차이</h3>{table}<p>S01 → S02 → S03 장소 검색 / S05 편집 / S04 기록 → V1 감상 · S04 → M2 → M3</p></section>'
    page='<!doctype html><html lang="ko"><meta charset="utf-8"><title>At that place UI 설계 · 프로토타입 반영</title><style>'+css+'</style><aside><b>UI 설계 · 화면 출처별 표시</b><p><a href="../00 설계서 안내.html">← 설계서 안내</a></p>'+''.join(nav)+'</aside><main><header><h1>At that place · UI 설계</h1><p>2026-10-26 · feature/ui-prototype 기준 화면 교체</p><a href="05 UI 설계.pptx">PowerPoint</a> · <a href="05 화면 명세.xlsx">화면 명세 Excel</a> · <a href="05 UI 설계.pdf">PDF</a></header>'+intro+''.join(articles)+'</main></html>'
    (ROOT/'UI/05 UI 화면 갤러리.html').write_text(page)
    result=dict(source_commit=commit,main_screens=len(SCREENS),replaced_ui=list(primary),
                captures=len(shots),supplemental_screens=len(shots)-len(primary),
                gallery_pages=page_ids,pdf_intro_pages=1,pdf_pages=len(page_ids)+1,pptx_slides=len(prs.slides))
    (ROOT/'UI/화면 반영.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False))


def update_index():
    """전체 설계서 재생성 시에도 UI 안내가 예전 편집용 도형 설명으로 돌아가지 않게 한다."""
    p=ROOT/'00 설계서 안내.html'
    s=p.read_text()
    s=s.replace('UI 화면과 APK 캡처 함께 보기 →','최신 프로토타입 UI 화면 보기 →')
    s=s.replace('편집 가능한 UI 원본','UI 설계와 프로토타입 화면')
    s=s.replace('글자·버튼·필드·배경을 개별 도형으로 수정합니다. 실제 APK 캡처는 별도 근거 이미지입니다.',
                'feature/ui-prototype의 7개 대응 UI를 실행 캡처로 교체하고 관련 상태 12개를 추가했습니다. 캡처는 앱 소스에서, 설명은 PowerPoint에서 편집합니다.')
    s=s.replace('SVG 원본은 각 갤러리 화면에서 엽니다.','화면별 출처·실행 경로·제품 계약 차이는 갤러리와 Excel에서 확인합니다.')
    s=s.replace('SVG는 벡터 편집기로 열어 수정합니다.','명세 기반 설계안의 SVG는 벡터 편집기로 수정합니다. 프로토타입 대응 SVG는 최신 화면을 담은 이미지 래퍼입니다.')
    if '2026-10-26 UI 화면 반영' not in s:
        s=s.replace('<h2>개발 설계서</h2>','<div class="info"><b>2026-10-26 UI 화면 반영</b><p>원격 feature/ui-prototype @ 99ab363 기준. 일정 요약·날짜별 장소·장소 검색·기록 보기·순서 편집과 관련 팝업을 반영했습니다. 웹 실행 캡처 19개, 기존 UI ID 33개와 제품 계약을 유지합니다. <a href="UI/프로토타입/출처.json">출처와 캡처 경로</a></p></div><h2>개발 설계서</h2>')
    p.write_text(s)
