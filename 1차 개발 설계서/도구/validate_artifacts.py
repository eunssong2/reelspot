"""생성된 설계 산출물의 연결·파일무결성·편집가능구조 검사. 앱 기능 테스트가 아니다."""
from pathlib import Path
import json,hashlib,re,zipfile,sys
from urllib.parse import unquote,urlsplit
from html.parser import HTMLParser
from xml.etree import ElementTree as ET
from docx import Document
from pptx import Presentation
import openpyxl,pymupdf
from openapi_spec_validator import validate
from design_model import ROOT,SCREENS,APIS,QA
from physical_model import TABLES as PHYSICAL_TABLES

class Links(HTMLParser):
    def __init__(self):super().__init__();self.links=[]
    def handle_starttag(self,tag,attrs):
        for k,v in attrs:
            if k in ['src','href']:self.links.append(v)

def main():
    checks=[];errors=[]
    def check(name,ok,detail):
        checks.append({'check':name,'passed':bool(ok),'detail':detail})
        if not ok:errors.append(name)
    spec=json.loads((ROOT/'부록/openapi.json').read_text());validate(spec)
    ops=[v['operationId'] for path in spec['paths'].values() for v in path.values()]
    check('OpenAPI3.1 스키마와 operation 고유성',len(ops)==len(set(ops))==len(APIS),len(ops))
    required={f'F{i:02}' for i in range(1,29)}
    check('28개 기능의 UI 연결',required<=set(f for s in SCREENS for f in s['features']),len(required))
    api_features=set(f for a in APIS for f in a['features'])
    check('서버 기능의 API 연결과 로컬 촬영 경계',(required-{'F12'})<=api_features,{'server_features':len(required-{'F12'}),'F12':'Android 카메라 로컬 처리. 후보 사용 뒤 F14/F15 접수·전송 연결'})
    check('화면 API 참조',all(set(s['api'])<=set(ops) for s in SCREENS),len(SCREENS))
    basis=json.loads((ROOT/'기준/기준선.json').read_text())
    check('패키지 내부 기능명세 SHA256',hashlib.sha256((ROOT/'기준/기능명세서.md').read_bytes()).hexdigest()==basis['sha256'],basis['sha256'])
    op=spec['components']['schemas']['Operation']
    check('업무 결과 키 operationId 통일','operationId' in op['required'] and 'id' not in op['properties'],op['required'])
    wire=spec['paths']['/upload-tasks/{taskId}/content']
    check('HEAD/PATCH 전송 계약', 'requestBody' not in wire['head'] and '204' in wire['head']['responses'] and 'application/offset+octet-stream' in wire['patch']['requestBody']['content'],list(wire))
    contract=json.loads((ROOT/'부록/계약예제/검증결과.json').read_text())
    canonical=hashlib.sha256(json.dumps(spec,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    check('현재 OpenAPI 예제 검증 결과',contract['passed'] and contract['openapi_canonical_sha256']==canonical,len(contract['checks']))
    db=json.loads((ROOT/'부록/DB/검증결과.json').read_text())
    check('현재 DDL의 실제 DB 검증 결과',db['passed'] and db['ddl_sha256']==hashlib.sha256((ROOT/'부록/DB/001_initial.sql').read_bytes()).hexdigest(),{'checks':len(db['checks']),'database':db['database']})
    actual=json.loads((ROOT/'부록/DB/실제 컬럼 사전.json').read_text())
    actual_set={(c['table_name'],c['column_name']) for c in actual}
    expected_set={(n,c[0]) for n,t in PHYSICAL_TABLES.items() for c in t['columns']}
    check('컬럼 사전과 실제 DB 컬럼 일치',actual_set==expected_set,{'tables':len(PHYSICAL_TABLES),'columns':len(actual_set)})
    isolated=json.loads((ROOT/'부록/독립 패키지 검증.json').read_text())
    check('상위 원본과 DayLog 없는 패키지 재생성',isolated['passed'],isolated)
    check('112개 기능 검수 시나리오',set(QA)==required and sum(len(x) for x in QA.values())==112,112)
    svg_paths=list((ROOT/'UI/SVG').glob('*.svg'))
    svg_ok=True
    for p in svg_paths:
        el=ET.parse(p)
        if any(x.tag.endswith('image') for x in el.iter()):svg_ok=False
        if not any(x.tag.endswith('text') for x in el.iter()):svg_ok=False
    check(f'{len(SCREENS)}개 SVG의 독립 텍스트/도형',len(svg_paths)==len(SCREENS) and svg_ok,len(svg_paths))
    issues=json.loads((ROOT/'부록/SVG 레이아웃 검사.json').read_text())
    check('SVG 텍스트 화면경계',len(issues)==0,issues)
    prs=Presentation(ROOT/'UI/05 UI 설계.pptx');layout_issues=[];native=[]
    for slide in prs.slides:
        texts='\n'.join(sh.text for sh in slide.shapes if sh.has_text_frame and sh.text.strip())
        for s in SCREENS:
            if texts.startswith(s['id']+'  '):native.append((s['id'],sum(1 for sh in slide.shapes if sh.has_text_frame)))
        for sh in slide.shapes:
            if min(sh.left,sh.top)<-1000 or sh.left+sh.width>prs.slide_width+1000 or sh.top+sh.height>prs.slide_height+1000:
                layout_issues.append({'slide':len(layout_issues),'shape':sh.name})
    check(f'PPTX {len(SCREENS)}개 화면의 네이티브 텍스트',len(native)==len(SCREENS) and all(n>=10 for _,n in native),native)
    check('PPTX 도형 슬라이드 경계',not layout_issues,layout_issues)
    documents=list(ROOT.glob('*.docx'))+list((ROOT/'근거').glob('*.docx'))
    for p in documents:
        with zipfile.ZipFile(p) as z:assert z.testzip() is None
        d=Document(p);check('Word 열기 '+p.stem,len(d.paragraphs)>8,{'paragraphs':len(d.paragraphs),'tables':len(d.tables)})
    workbooks=list(ROOT.glob('*.xlsx'))+list((ROOT/'UI').glob('*.xlsx'))+list((ROOT/'부록/DB').glob('*.xlsx'))
    for p in workbooks:
        w=openpyxl.load_workbook(p,data_only=False);check('Excel 열기 '+p.stem,len(w.sheetnames)>2,w.sheetnames)
    w=openpyxl.load_workbook(ROOT/'06 검수 및 추적.xlsx')
    check('기능검수 결과 미실행 유지',all(w['기능검수'].cell(r,9).value=='미실행' for r in range(2,114)),112)
    broken=[]
    for p in ROOT.rglob('*.html'):
        if '.venv' in p.parts or '__pycache__' in p.parts:continue
        parser=Links();parser.feed(p.read_text())
        for link in parser.links:
            u=urlsplit(link)
            if u.scheme or not u.path:continue
            target=(p.parent/unquote(u.path)).resolve()
            if not target.exists():broken.append({'from':str(p.relative_to(ROOT)),'link':link})
    check('로컬 HTML 링크/이미지',not broken,broken)
    captures=list((ROOT/'근거/APK').glob('*.png'));capture_errors=[]
    for p in captures:
        m=json.loads(p.with_suffix('.json').read_text())
        if hashlib.sha256(p.read_bytes()).hexdigest()!=m['sha256']:capture_errors.append(p.name)
        ET.parse(p.with_suffix('.xml'))
    check('APK 캡처/메타/XML 무결성',not capture_errors,{'count':len(captures),'errors':capture_errors})
    pdfs=list(ROOT.glob('*.pdf'))+list((ROOT/'UI').glob('*.pdf'))+list((ROOT/'근거').glob('*.pdf'))
    for p in pdfs:
        pdf=pymupdf.open(p);check('PDF 열기 '+p.stem,len(pdf)>0 and bool(pdf[0].get_text().strip()),len(pdf))
        if p.stem=='05 UI 설계':check('UI PDF 화면수',len(pdf)==len(SCREENS),len(pdf))
    result={'scope':'설계 산출물 검증. 실제 앱 F01~F28 구현 검수는 미실행. PPTX는 구조 검사, 대응 SVG와 PDF는 시각 검토.','passed':not errors,'checks':checks,'errors':errors}
    (ROOT/'부록/산출물 검증.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    manifest=[]
    for p in sorted(ROOT.rglob('*')):
        if not p.is_file() or '__pycache__' in p.parts or '.venv' in p.parts or p.name=='파일 목록 SHA256.json':continue
        manifest.append({'path':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    (ROOT/'부록/파일 목록 SHA256.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print(json.dumps({'passed':not errors,'checks':len(checks),'files':len(manifest),'errors':errors},ensure_ascii=False))
    if errors:sys.exit(1)
if __name__=='__main__':main()
