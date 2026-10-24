import json,hashlib,subprocess
from pathlib import Path
from PIL import Image
import openpyxl
from design_model import ROOT,DATE
from build_design import Doc

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def main():
    apk=ROOT.parent/'기능 명세서 자료/ReelSpot-Demo.apk'
    if not apk.exists():
        assert (ROOT/'근거/APK/실행검토.json').exists(),'APK도 기존 실행 근거도 없습니다.'
        print('외부 APK가 없어 패키지에 보관된 실제 실행 근거를 유지합니다. 새 실행 결과를 생성하지 않았습니다.')
        return
    captures=[]
    for p in sorted((ROOT/'근거/APK').glob('*.png')):
        m=json.loads(p.with_suffix('.json').read_text())
        assert sha(p)==m['sha256']
        assert Image.open(p).size==(1080,2400)
        captures.append({k:m[k] for k in ['id','captured_at','action','sha256']})
    metadata={'review_date':DATE,'apk':str(apk.relative_to(ROOT.parent)),
      'apk_sha256':sha(apk),'package':'com.reelspot.app.demo','version_name':'0.1.0','version_code':1,
      'environment':{'avd':'daylog35','android':'15','api':35,'abi':'x86_64','resolution':'1080x2400','mode':'no-window Android emulator, ADB screen capture'},
      'launch_command':'adb -s emulator-5554 shell am start --user 0 -n com.reelspot.app.demo/.MainActivity',
      'observed':['일정 선택','서울 여행 개요·날짜 선택','날짜별 장소 목록','장소별 사진/영상 카드·작성자·시각·설명','업로드 후보 시트','기기 가져오기·사진/영상 촬영 메뉴','예시 지도 카페 분류·결과','날짜 제목·방문 시각·순서 편집'],
      'verified_mutations':[{'action':'서울숲 산책길 아래로 이동 후 변경 저장','before':'서울숲 산책길→오후의 커피','after':'오후의 커피→서울숲 산책길','evidence':['09-plan','10-order-edited','11-order-saved']},
         {'action':'예시 결과의 한남 브런치 카페 추가','before':'10월17일 장소4개','after':'장소5개·한남 브런치 카페·시간미정','evidence':['13-add-candidate','14-place-added']}],
      'not_verified':['실제 카카오/Google 인증','실제 서버 파일 업로드와 여러 기기 동기화','실제 카카오 지도 검색','카메라 실제 파일 생성·저장','서버 위치 제거·콘텐츠 검사','보관 만료·운영·종료 후 권한','Android12 최저 실기기·성능목표'],
      'observed_differences':['업로드 시트 최대10개→최신명세120개','날짜 화면을 경유→최신명세일정전체직접상세','날짜 전체 메모→최신명세제외'],
      'captures':captures}
    (ROOT/'근거/APK/실행검토.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2))
    book=openpyxl.load_workbook(ROOT.parent/'설계 자료/화면 명세 데이터.xlsx',data_only=True)
    inventory=[]
    for rel in ['기능명세서.md','설계 자료/프로젝트 소개.pdf','설계 자료/화면 명세 데이터.xlsx','기능 명세서 자료/ReelSpot-Demo.apk']:
        p=ROOT.parent/rel;inventory.append({'path':rel,'bytes':p.stat().st_size,'sha256':sha(p)})
    (ROOT/'근거/자료/자료 검토.json').write_text(json.dumps({'sources':inventory,'pdf_pages':14,'pdf_review':'텍스트 추출에서 한글 누락을 확인하여 14페이지를 렌더링해 시각 검토','workbook_sheets':[{'sheet':s.title,'rows':s.max_row,'columns':s.max_column,'embedded_images':len(s._images)} for s in book]},ensure_ascii=False,indent=2))
    d=Doc('08 APK 실행 검토','ReelSpot 데모 APK 실행과 UI 관찰','실제 실행 근거 · 1차 신규 설계/실서비스 검수와 구분')
    d.h('1 실행 환경과 식별')
    d.table(['항목','값'],[['기준일',DATE],['APK',metadata['apk']],['SHA256',metadata['apk_sha256']],['패키지·버전','com.reelspot.app.demo · 0.1.0 (code1)'],['에뮬레이터','daylog35 · Android15(API35) · x86_64 · 1080×2400'],['실행방식','디스플레이 연결 없이 Android 에뮬레이터를 부팅하고 ADB로 앱 실행·터치·원본 화면 캡처'],['실행명령',metadata['launch_command']]])
    d.p('APK를 직접 설치·실행해 예시 데이터 모드로 화면을 탐색했다. 이전 artifacts/reelspot-review 캡처를 이번 실행 근거로 재사용하지 않았다. 부팅 중 시스템 UI 대기 창이 나타났고 부팅·사용자 잠금 해제 뒤 앱 실행을 확인했다. 이 현상을 앱 기능 결함으로 판정하지 않았다.')
    d.h('2 관찰과 실제 변경 확인')
    for s in metadata['observed']:d.p(s)
    d.table(['실행한 변경','변경 전','저장·추가 뒤','근거'],[[m['action'],m['before'],m['after'],' → '.join(m['evidence'])] for m in metadata['verified_mutations']])
    d.h('3 최신 설계에 반영한 차이')
    d.table(['데모 관찰','최신 명세','반영'],[['여행→날짜→장소 탐색','모임→일정 전체 상세, 날짜는 필터','UI05·UI08·UI09로 정리'],['업로드 후보 최대10개 안내','촬영/기기 혼합1~120개','UI14·UI16·서버검증 변경'],['날짜 전체 장소 메모 입력','날짜 제목만, 공통 메모 제외','UI12에서 메모 제거'],['카드별 작성자·시간·설명과 정렬','파일당 카드·묶음 정렬·선택 저장','UI09·UI13·UI18·UI19 참고'],['지도에 예시 검색/실제 지도 선택','국내 카카오맵 실제 연동','예시 결과를 실지도 성공으로 보고하지 않음']])
    d.h('4 확인 범위의 한계')
    d.p('실제 데모 UI 관찰과 예시 데이터 변경을 확인한 결과다. 신규 화면의 기능이나 서버의 완료 조건을 검증한 결과는 아니다.')
    for s in metadata['not_verified']:d.p(s+' — 이번 실행에서 검증하지 않음')
    d.h('5 캡처 목록')
    d.table(['ID','행동','시각'],[[m['id'],m['action'],m['captured_at']] for m in captures])
    d.h('6 실제 화면')
    for i,m in enumerate(captures):
        d.h(m['id']+'  '+m['action'],2)
        p=ROOT/'근거/APK'/(m['id']+'.png')
        d.doc.add_picture(str(p),width=__import__('docx').shared.Inches(2.1))
        d.parts.append(f'<img style="width:220px;max-height:510px;object-fit:contain" src="../../근거/APK/{p.name}">')
        d.p('근거 파일: '+m['id']+'.png / .xml / .json')
    d.save();(ROOT/'08 APK 실행 검토.docx').replace(ROOT/'근거/08 APK 실행 검토.docx')
    print('captured screens',len(captures),'verified demo mutations',len(metadata['verified_mutations']))
if __name__=='__main__':main()
